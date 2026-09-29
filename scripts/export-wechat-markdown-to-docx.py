#!/usr/bin/env python3
"""Export a WeChat Markdown article to a clean, image-complete Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, Iterable, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


BODY_FONT = "Hiragino Sans GB"
MONO_FONT = "Menlo"
BLACK = "20252B"
GRAY = "595959"
LIGHT_GRAY = "888888"
ORANGE = "E87522"
ORANGE_DARK = "C85D12"
ORANGE_PALE = "FFF3E6"
ORANGE_WASH = "FFF9F2"
PALE_BORDER = "F2C6A0"


def parse_front_matter(text: str) -> Tuple[Dict[str, str], str]:
    if not text.startswith("---\n"):
        return {}, text
    end = text.find("\n---\n", 4)
    if end == -1:
        return {}, text
    metadata: Dict[str, str] = {}
    for raw in text[4:end].splitlines():
        if ":" not in raw:
            continue
        key, value = raw.split(":", 1)
        metadata[key.strip()] = value.strip().strip('"').strip("'")
    return metadata, text[end + 5 :]


def set_run_font(run, name: str = BODY_FONT, size: Optional[Pt] = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_shading(run, fill: str) -> None:
    run_pr = run._element.get_or_add_rPr()
    shading = run_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        run_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color: str, size: int = 18, space: int = 7) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def add_hyperlink(paragraph, text: str, url: str, size: Pt = Pt(11.25)):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ORANGE_DARK)
    run_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    run_pr.append(fonts)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(size.pt * 2)))
    run_pr.append(font_size)
    font_size_cs = OxmlElement("w:szCs")
    font_size_cs.set(qn("w:val"), str(int(size.pt * 2)))
    run_pr.append(font_size_cs)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"\[([^\]]+)\]\((https?://[^)]+)\)|"
    r"\*\*([^*]+)\*\*|"
    r"`([^`]+)`|"
    r"(?<!\*)\*([^*]+)\*(?!\*)"
)


def add_inline_markdown(paragraph, text: str, base_size: Pt = Pt(11.25)) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=base_size)
        if match.group(1) is not None:
            add_hyperlink(paragraph, match.group(1), match.group(2), base_size)
        elif match.group(3) is not None:
            run = paragraph.add_run(match.group(3))
            run.bold = True
            set_run_font(run, size=base_size)
        elif match.group(4) is not None:
            run = paragraph.add_run(match.group(4))
            set_run_font(run, MONO_FONT, Pt(max(base_size.pt - 0.75, 9.5)))
            run.font.color.rgb = RGBColor.from_string(ORANGE_DARK)
            set_run_shading(run, ORANGE_PALE)
        elif match.group(5) is not None:
            run = paragraph.add_run(match.group(5))
            run.italic = True
            set_run_font(run, size=base_size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size)


def style_paragraph(paragraph, *, after=9, before=0, line=1.75, first_line=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        fmt.first_line_indent = Pt(22)


def configure_document(document: Document, metadata: Dict[str, str]) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11.25)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = document.styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 20, 8),
        ("Heading 2", 15, 18, 7),
        ("Heading 3", 12.5, 12, 5),
    ):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(
            ORANGE_DARK if name in {"Heading 1", "Heading 2"} else BLACK
        )
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Article Author" not in document.styles:
        author_style = document.styles.add_style("Article Author", WD_STYLE_TYPE.PARAGRAPH)
        author_style.font.name = BODY_FONT
        author_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        author_style.font.size = Pt(10.5)
        author_style.font.color.rgb = RGBColor.from_string(ORANGE_DARK)
        author_style.paragraph_format.space_after = Pt(7)

    if "Article Digest" not in document.styles:
        digest_style = document.styles.add_style("Article Digest", WD_STYLE_TYPE.PARAGRAPH)
        digest_style.font.name = BODY_FONT
        digest_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        digest_style.font.size = Pt(10.5)
        digest_style.font.italic = False
        digest_style.font.color.rgb = RGBColor.from_string(GRAY)
        digest_style.paragraph_format.line_spacing = 1.55
        digest_style.paragraph_format.space_after = Pt(10)

    props = document.core_properties
    props.title = metadata.get("title", "")
    props.author = metadata.get("author", "")
    props.subject = metadata.get("digest", "")


def image_size(path: Path, max_width: float = 6.55, max_height: float = 5.15) -> Tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def set_picture_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_image(
    document: Document,
    source: Path,
    relative_path: str,
    alt: str,
    *,
    keep_with_next: bool,
) -> None:
    image_path = (source.parent / relative_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(f"Image not found: {image_path}")
    width, height = image_size(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    set_picture_alt_text(shape, alt)


def add_code_block(document: Document, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.keep_together = True
    set_paragraph_shading(paragraph, ORANGE_PALE)
    set_paragraph_left_border(paragraph, ORANGE, size=20, space=8)
    code = "\n".join(lines) or " "
    run = paragraph.add_run(code)
    set_run_font(run, MONO_FONT, Pt(9.75))
    run.font.color.rgb = RGBColor.from_string(BLACK)


def markdown_to_docx(
    source: Path,
    output: Path,
    *,
    skip_first_image: bool = False,
    publish_opening: bool = False,
    display_title: str | None = None,
    display_author: str | None = None,
) -> None:
    metadata, body = parse_front_matter(source.read_text(encoding="utf-8"))
    document_metadata = dict(metadata)
    if display_title:
        document_metadata["title"] = display_title
    if display_author:
        document_metadata["author"] = display_author
    document = Document()
    configure_document(document, document_metadata)

    lines = body.splitlines()
    index = 0
    title_seen = False
    image_count = 0
    keep_next_body = False
    in_references = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if publish_opening and stripped.startswith(("项目入口：", "项目入口:")):
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip("\n"))
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_count += 1
            if skip_first_image and image_count == 1:
                index += 1
                continue
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            next_is_caption = (
                next_index < len(lines)
                and bool(re.fullmatch(r"\*[^*]+\*", lines[next_index].strip()))
            )
            add_image(
                document,
                source,
                image_match.group(2),
                image_match.group(1),
                keep_with_next=next_is_caption,
            )
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            in_references = heading_text == "参考资料"
            if level == 1 and not title_seen:
                visible_title = display_title or heading_text
                title_size = Pt(18.5 if publish_opening else 24)
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_markdown(paragraph, visible_title, title_size)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(BLACK)
                if metadata.get("author"):
                    author = document.add_paragraph(style="Article Author")
                    byline_label = "原创" if publish_opening else "作者"
                    byline_name = display_author or metadata["author"]
                    author_run = author.add_run(f"{byline_label}：{byline_name}")
                    set_run_font(author_run, size=Pt(10.5))
                    author_run.font.color.rgb = RGBColor.from_string(
                        LIGHT_GRAY if publish_opening else ORANGE_DARK
                    )
                if metadata.get("digest") and not publish_opening:
                    digest = document.add_paragraph(style="Article Digest")
                    digest_run = digest.add_run(metadata["digest"])
                    set_run_font(digest_run, size=Pt(10.5))
                    digest_run.font.color.rgb = RGBColor.from_string(GRAY)
                title_seen = True
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
                add_inline_markdown(paragraph, heading_text, paragraph.style.font.size or Pt(14))
                heading_color = ORANGE_DARK if level <= 2 else BLACK
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(heading_color)
            keep_next_body = level == 3 and heading_text == "原理"
            index += 1
            continue

        ordered_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.65
            number_run = paragraph.add_run(f"{ordered_match.group(1)}. ")
            set_run_font(number_run, size=Pt(11.25))
            number_run.bold = True
            number_run.font.color.rgb = RGBColor.from_string(ORANGE_DARK)
            add_inline_markdown(paragraph, ordered_match.group(2))
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(2 if in_references else 5)
            paragraph.paragraph_format.line_spacing = 1.35 if in_references else 1.65
            bullet_run = paragraph.add_run("• ")
            set_run_font(bullet_run, size=Pt(9.5) if in_references else Pt(11.25))
            bullet_run.bold = True
            bullet_run.font.color.rgb = RGBColor.from_string(ORANGE)
            add_inline_markdown(
                paragraph,
                bullet_match.group(1),
                Pt(9.5) if in_references else Pt(11.25),
            )
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.42)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.55
            set_paragraph_shading(paragraph, ORANGE_WASH)
            set_paragraph_left_border(paragraph, ORANGE, size=16, space=7)
            add_inline_markdown(paragraph, stripped[2:], Pt(10.75))
            for run in paragraph.runs:
                if run.font.name != MONO_FONT:
                    run.font.color.rgb = RGBColor.from_string(GRAY)
            index += 1
            continue

        paragraph = document.add_paragraph()
        is_caption = bool(re.fullmatch(r"\*[^*]+\*", stripped))
        if is_caption:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_with_next = False
            add_inline_markdown(paragraph, stripped, Pt(9.5))
            for run in paragraph.runs:
                run.italic = False
                run.font.color.rgb = RGBColor.from_string(LIGHT_GRAY)
        else:
            style_paragraph(paragraph)
            if keep_next_body:
                paragraph.paragraph_format.keep_with_next = True
                keep_next_body = False
            add_inline_markdown(paragraph, stripped)
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--skip-first-image",
        action="store_true",
        help="Omit the first body image when the publishing platform uses a separate cover upload.",
    )
    parser.add_argument(
        "--publish-opening",
        action="store_true",
        help="Use a clean WeChat opening with an original-byline and no digest or project-intro paragraph.",
    )
    parser.add_argument(
        "--display-title",
        help="Override the visible document title without changing the Markdown source.",
    )
    parser.add_argument(
        "--display-author",
        help="Override the visible byline without changing the Markdown source.",
    )
    args = parser.parse_args()
    markdown_to_docx(
        args.source.resolve(),
        args.output.resolve(),
        skip_first_image=args.skip_first_image,
        publish_opening=args.publish_opening,
        display_title=args.display_title,
        display_author=args.display_author,
    )


if __name__ == "__main__":
    main()
